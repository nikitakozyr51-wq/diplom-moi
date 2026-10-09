#!/usr/bin/env python3
"""Export the working plan to DOCX using the university's basic text settings."""
from pathlib import Path
import re
import subprocess

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "plan.docx"
BUILD = ROOT / "build"
BUILD.mkdir(exist_ok=True)

reference = Document()
section = reference.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
section.different_first_page_header_footer = True

for style_name in ["Normal", "Body Text", "First Paragraph", "Title", "Subtitle",
                   "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Caption"]:
    if style_name not in reference.styles:
        style = reference.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = reference.styles["Normal"]
    else:
        style = reference.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12 if style_name == "Caption" else 14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    style.font.bold = style_name.startswith("Heading") or style_name == "Title"
    style.font.italic = False
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if style_name.startswith("Heading") or style_name in ["Title", "Subtitle"]:
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
        pf.space_before, pf.space_after = Pt(12), Pt(6)
    if style_name == "Heading 2":
        pf.page_break_before = True
    if style_name == "Heading 4":
        pf.space_before, pf.space_after = Pt(6), Pt(3)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
run = footer.add_run('2')
run.font.name, run.font.size = "Times New Roman", Pt(12)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.remove(run._r)
field.append(run._r)
footer._p.append(field)
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
reference.settings._element.append(update)
reference.save(BUILD / "reference.docx")

text = (ROOT / "plan.md").read_text(encoding="utf-8")
# Promote the actual thesis subsection labels into Word heading styles.
text = re.sub(r"(?m)^\*\*(\d\.\d [^*]+)\*\*\s*", r"#### \1\n\n", text)
(BUILD / "plan-export.md").write_text(text, encoding="utf-8")
subprocess.run(["pandoc", str(BUILD / "plan-export.md"),
                "--reference-doc", str(BUILD / "reference.docx"),
                "--from", "markdown", "--to", "docx", "-o", str(OUTPUT)], check=True)

doc = Document(OUTPUT)
# Match the 12 pt / single-spaced table requirement after conversion.
for table in doc.tables:
    table.style = 'Table Grid'
    table.autofit = True
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles['Normal']
                pf = paragraph.paragraph_format
                pf.line_spacing = 1
                pf.first_line_indent = Cm(0)
                pf.space_before = pf.space_after = Pt(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
                    run.font.bold = False
                    run.font.color.rgb = RGBColor(0, 0, 0)
    # Repeat header rows when Word paginates a table.
    properties = table.rows[0]._tr.get_or_add_trPr()
    for old in list(properties.findall(qn('w:tblHeader'))):
        properties.remove(old)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)

for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading") and paragraph.runs:
        paragraph.runs[-1].text = paragraph.runs[-1].text.rstrip().rstrip(".")
    if paragraph.style.name == "Heading 2":
        for run in paragraph.runs:
            run.text = run.text.upper()
    if paragraph.style.name in ["Normal", "Body Text", "First Paragraph"]:
        for run in paragraph.runs:
            run.font.bold = False
doc.save(OUTPUT)
print(OUTPUT)
