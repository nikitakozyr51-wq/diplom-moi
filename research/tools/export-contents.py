#!/usr/bin/env python3
"""Export the short proposed thesis contents for discussion with the supervisor."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
section.top_margin = section.bottom_margin = Cm(2)
section.footer_distance = Cm(1)
normal = document.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.widow_control = True

text = (ROOT / 'contents-plan.md').read_text()
in_contents = False
for block in text.strip().split('\n\n'):
    block = block.strip()
    if not block:
        continue
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if block.startswith('# '):
        paragraph.add_run(block[2:].upper()).bold = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(8)
    elif block.startswith('## '):
        paragraph.add_run(block[3:]).bold = True
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_before = Pt(6)
    else:
        paragraph.add_run(block)
        paragraph.paragraph_format.keep_together = True
        if re.match(r'^\d+\.\d+ ', block):
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)
        elif block == 'Перечень сокращений':
            in_contents = True
            paragraph.paragraph_format.space_before = Pt(6)
        elif not in_contents:
            paragraph.paragraph_format.space_after = Pt(6)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run = footer.add_run('1')
run.font.name, run.font.size = 'Times New Roman', Pt(12)
footer._p.remove(run._r)
field.append(run._r)
footer._p.append(field)
document.save(ROOT / 'contents-plan.docx')
print(ROOT / 'contents-plan.docx')
