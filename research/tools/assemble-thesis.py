#!/usr/bin/env python3
"""Combine the official cover, complete planned contents and written thesis.

The original template remains untouched. Page numbers are refreshed from a
rendered PDF using --pagination-pdf. Future sections have no page references.
"""
import argparse
from copy import deepcopy
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.packuri import PackURI
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TOPIC = ('Исследование и анализ применения технологий автономного судовождения '
         '(MASS) при плавании в узкостях и перспективы их внедрения на морском флоте')
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--body', type=Path, default=ROOT / 'build/draft-body.docx')
parser.add_argument('--output', type=Path, default=ROOT / 'drafts/working-draft.docx')
parser.add_argument('--pagination-pdf', type=Path)
args = parser.parse_args()
body = Document(args.body)
doc = Document(ROOT / 'requirements/originals/title-template.docx')
cover = list(doc.paragraphs)

# Cover typography belongs to the supplied official form, rather than to the
# 14 pt / 1.5-spaced manuscript. Preserve the emblem and existing run settings.
title_style = doc.styles.add_style('VKR Title Page', WD_STYLE_TYPE.PARAGRAPH)
title_style.font.name, title_style.font.size = 'Times New Roman', Pt(12)
title_style.paragraph_format.line_spacing = 1.15
title_style.paragraph_format.space_before = title_style.paragraph_format.space_after = Pt(0)
title_style.paragraph_format.first_line_indent = Cm(0)
for paragraph in cover:
    paragraph.style = title_style

def fill(paragraph, text, size=12):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name, run.font.size = 'Times New Roman', Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = False

cover[0]._element.getparent().remove(cover[0]._element)  # Appendix label of blank form.
fill(cover[10], 'Направление подготовки (специальность) 26.05.05 Судовождение')
fill(cover[24], 'Обучающегося Козырева Никиты Алексеевича, группа С-612')
fill(cover[26], 'Вид работы: выпускная квалификационная работа специалиста')
fill(cover[31], 'Тема: ' + TOPIC, 14)
cover[32]._element.getparent().remove(cover[32]._element)
fill(cover[45], '2026')

# Import manuscript styles and numbering. The title keeps its separate style.
for style in body.styles:
    if style.style_id == title_style.style_id:
        continue
    for previous in list(doc.styles._element):
        name = previous.find(qn('w:name'))
        if (previous.get(qn('w:styleId')) == style.style_id or
                (name is not None and name.get(qn('w:val'), '').casefold() == style.name.casefold())):
            doc.styles._element.remove(previous)
    doc.styles._element.append(deepcopy(style._element))
numbering = NumberingPart(PackURI('/word/numbering.xml'), CT.WML_NUMBERING,
                          deepcopy(body.part.numbering_part._element), doc.part.package)
doc.part.relate_to(numbering, RT.NUMBERING)

def style(name, *, bold=False, center=False, outline=None):
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles['Normal']
    s.font.name, s.font.size = 'Times New Roman', Pt(14)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0, 0, 0)
    f = s.paragraph_format
    f.first_line_indent = Cm(0)
    f.space_before, f.space_after = Pt(0), Pt(0)
    f.line_spacing = 1.5
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    if outline is not None:
        level = OxmlElement('w:outlineLvl')
        level.set(qn('w:val'), str(outline))
        s._element.get_or_add_pPr().append(level)
    return s

front = style('VKR Front Heading', bold=True, center=True, outline=9)
front.paragraph_format.page_break_before = True
front.paragraph_format.keep_with_next = True
front.paragraph_format.space_after = Pt(14)
structural = style('VKR Structural Heading', bold=True, center=True, outline=0)
structural.paragraph_format.page_break_before = True
structural.paragraph_format.keep_with_next = True
structural.paragraph_format.space_after = Pt(14)
toc_style = style('VKR TOC Entry', outline=9)
toc_style.paragraph_format.keep_together = True
toc_style.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
abbreviation_style = style('VKR Abbreviation', outline=9)
abbreviation_style.paragraph_format.tab_stops.add_tab_stop(Cm(3))
abbreviation_style.paragraph_format.left_indent = Cm(3)
abbreviation_style.paragraph_format.first_line_indent = Cm(-3)

doc.add_paragraph('СОДЕРЖАНИЕ', front)
entries = []
glossary = doc.add_paragraph('ПЕРЕЧЕНЬ СОКРАЩЕНИЙ', structural)
abbreviations = [
    ('ММО', 'Международная морская организация'),
    ('MASS', 'Maritime Autonomous Surface Ship — морское автономное надводное судно'),
    ('МППСС-72', 'Международные правила предупреждения столкновений судов в море 1972 года'),
    ('ГНСС', 'глобальная навигационная спутниковая система'),
    ('РЛС', 'радиолокационная станция'),
    ('АИС', 'автоматическая идентификационная система'),
]
for abbreviation, expansion in abbreviations:
    doc.add_paragraph(abbreviation + '\t' + expansion, abbreviation_style)

def bookmark(paragraph, number):
    name = 'VKRSection' + str(number)
    begin = OxmlElement('w:bookmarkStart')
    begin.set(qn('w:id'), str(1000 + number))
    begin.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(1000 + number))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, begin)
    paragraph._p.append(end)
    return name

entries.append({'heading': glossary.text, 'label': 'Перечень сокращений', 'bookmark': bookmark(glossary, 0), 'level': 1})
relationship_map = {}
for relationship in body.part.rels.values():
    if relationship.is_external:
        relationship_map[relationship.rId] = doc.part.relate_to(relationship.target_ref, relationship.reltype, is_external=True)
for element in body.element.body:
    if element.tag == qn('w:sectPr'):
        continue
    copied = deepcopy(element)
    for node in copied.iter():
        rid = node.get(qn('r:id'))
        if rid in relationship_map:
            node.set(qn('r:id'), relationship_map[rid])
    doc.element.body.insert(len(doc.element.body) - 1, copied)

started_body = False
for paragraph in doc.paragraphs:
    if paragraph.style.name not in ('Heading 1', 'Heading 2'):
        continue
    started_body = True
    if paragraph.style.name == 'Heading 1':
        paragraph.paragraph_format.page_break_before = True
        if paragraph.text in ('ВВЕДЕНИЕ', 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ'):
            paragraph.style = structural
    label = paragraph.text
    if not re.match(r'^\d', label):
        label = label[:1].upper() + label[1:].lower()
    elif paragraph.style.name != 'Heading 2':
        number, title = label.split(' ', 1)
        label = number + ' ' + title[:1].upper() + title[1:].lower()
    entries.append({'heading': paragraph.text, 'label': label,
                    'bookmark': bookmark(paragraph, len(entries)),
                    'level': 2 if paragraph.style.name == 'Heading 2' else 1})

normalize = lambda s: re.sub(r'\s+', '', s).casefold()
written = {normalize(e['label']): e for e in entries}
planned = []
for line in (ROOT / 'contents-plan.md').read_text().splitlines():
    label = line.removeprefix('## ').strip()
    if (re.match(r'^\d[. ]', label) or label in
            ('Перечень сокращений', 'Введение', 'Заключение',
             'Список использованных источников', 'Приложения (при необходимости)')):
        entry = written.pop(normalize(label), None)
        if entry is None:
            entry = {'heading': label, 'label': label, 'bookmark': None,
                     'level': 2 if re.match(r'^\d\.\d', label) else 1}
        entry['written'] = entry['bookmark'] is not None
        entry['page'] = None
        planned.append(entry)
if written:
    raise ValueError('Written headings absent from the approved plan: ' + str(list(written)))
entries = planned
if args.pagination_pdf:
    from pypdf import PdfReader
    raw_pages = [p.extract_text() for p in PdfReader(args.pagination_pdf).pages]
    pages = [normalize(p) for p in raw_pages]
    start = next(i for i, text in enumerate(raw_pages)
                 if any(line.strip() == 'ПЕРЕЧЕНЬ СОКРАЩЕНИЙ' for line in text.splitlines()))
    for entry in entries:
        if not entry['written']:
            continue
        matches = [i + 1 for i, text in enumerate(pages) if i >= start and normalize(entry['heading']) in text]
        if len(matches) != 1:
            raise ValueError('Cannot locate heading uniquely: ' + entry['heading'] + ': ' + str(matches))
        entry['page'] = matches[0]
else:
    for entry in entries:
        if entry['written']:
            entry['page'] = 0  # Refreshed before delivery, never published as zero.

def field_char(paragraph, kind):
    element = OxmlElement('w:fldChar')
    element.set(qn('w:fldCharType'), kind)
    paragraph.add_run()._r.append(element)

toc_heading = next(p for p in doc.paragraphs if p.text == 'СОДЕРЖАНИЕ')
last = toc_heading._p
for index, entry in enumerate(entries):
    paragraph = doc.add_paragraph(style=toc_style)
    last.addnext(paragraph._p)
    last = paragraph._p
    # Keep the approved outline stable when Word updates page fields.
    # An outer TOC field would discard headings that have no body yet.
    if entry['level'] == 2:
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.tab_stops.clear_all()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if not entry['written']:
        paragraph.add_run(entry['label'])
        continue
    # Reserve room so a long heading cannot touch its page number.
    paragraph.paragraph_format.right_indent = Cm(0.8)
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'), entry['bookmark'])
    run, txt = OxmlElement('w:r'), OxmlElement('w:t')
    txt.text = entry['label']
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run('\u00a0\t')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGEREF ' + entry['bookmark'] + ' \\h')
    run, txt = OxmlElement('w:r'), OxmlElement('w:t')
    txt.text = str(entry['page'])
    run.append(txt)
    field.append(run)
    paragraph._p.append(field)

section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
section.top_margin = section.bottom_margin = Cm(2)
section.header_distance = section.footer_distance = Cm(1)
section.different_first_page_header_footer = True
for container in (section.header, section.first_page_header, section.footer, section.first_page_footer):
    for element in list(container._element):
        container._element.remove(element)
footer = section.footer.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
footer.paragraph_format.line_spacing = 1
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run = footer.add_run('2')
run.font.name, run.font.size = 'Times New Roman', Pt(12)
footer._p.remove(run._r)
field.append(run._r)
footer._p.append(field)
for element in list(section._sectPr.findall(qn('w:pgNumType'))):
    section._sectPr.remove(element)
update = doc.settings._element.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    doc.settings._element.append(update)
update.set(qn('w:val'), 'true')
doc.save(args.output)
(ROOT / 'build/thesis-contents.json').write_text(json.dumps(entries, ensure_ascii=False, indent=2) + '\n')
print(args.output)
