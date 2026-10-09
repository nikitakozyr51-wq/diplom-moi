#!/usr/bin/env python3
"""Export the working draft or a selected chapter with the university's Word settings."""
from pathlib import Path
import argparse
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--input', type=Path, default=root / 'drafts' / 'working-draft.md')
parser.add_argument('--output', type=Path, default=root / 'drafts' / 'working-draft.docx')
args = parser.parse_args()
build = root / 'build'
reference = Document(build / 'reference.docx')
hyperlink = reference.styles['Hyperlink'] if 'Hyperlink' in reference.styles else reference.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER)
hyperlink.font.name = 'Times New Roman'
hyperlink.font.size = Pt(14)
hyperlink.font.color.rgb = RGBColor(0, 0, 0)
hyperlink.font.underline = False
reference.styles['Heading 1'].paragraph_format.page_break_before = True
reference.styles['Heading 2'].paragraph_format.page_break_before = False
for style_name in ('Normal', 'Body Text', 'First Paragraph', 'Heading 1', 'Heading 2'):
    properties = reference.styles[style_name]._element.get_or_add_rPr()
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'ru-RU')
    properties.append(language)
section = reference.sections[0]
# This is a chapter fragment, with no university title page.
section.different_first_page_header_footer = False
reference.save(build / 'draft-reference.docx')
assemble = args.input.resolve() == (root / 'drafts' / 'working-draft.md').resolve()
output = build / 'draft-body.docx' if assemble else args.output
text = args.input.read_text(encoding='utf-8')
body, marker, references = text.partition('# Список использованных источников')
if marker:
    # The university guide requires plain source numbers without a dot.
    references = re.sub(r'(?m)^(\d+)\. (.+)$', r'\1 \2\n', references)
    text = body + marker + references
export_input = build / (args.input.stem + '-word.md')
export_input.write_text(text, encoding='utf-8')
subprocess.run(['pandoc', str(export_input),
                '--reference-doc', str(build / 'draft-reference.docx'),
                '--from', 'markdown', '--to', 'docx', '-o', str(output)], check=True)
doc = Document(output)
first_chapter = True
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith('Heading'):
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 0, 0)
        if paragraph.style.name == 'Heading 1':
            if first_chapter:
                paragraph.paragraph_format.page_break_before = False
                first_chapter = False
            for run in paragraph.runs:
                run.text = run.text.upper()
    elif re.match(r'^Таблица \d+\.\d+', paragraph.text):
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = False
    else:
        for run in paragraph.runs:
            run.font.bold = False

for table_index, table in enumerate(doc.tables):
    table.style = 'Table Grid'
    first_header = table.cell(0, 0).text
    widths = ([2.3, 5.8, 8.4] if first_header.startswith('Степень') else
              [4.2, 5.8, 6.5] if first_header == 'Звено системы' else [5.0, 5.5, 6.0])
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            for paragraph in cell.paragraphs:
                # Current tables fit one page; keep them together with their
                # caption instead of emitting an unlabelled continuation.
                paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                paragraph.style = doc.styles['Normal']
                pf = paragraph.paragraph_format
                pf.line_spacing = 1
                pf.first_line_indent = Cm(0)
                pf.space_before = pf.space_after = Pt(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    props = table.rows[0]._tr.get_or_add_trPr()
    for existing in list(props.findall(qn('w:tblHeader'))):
        props.remove(existing)
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    props.append(repeat)
    for row in table.rows:
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
doc.save(output)
if assemble:
    subprocess.run([sys.executable, str(root / 'tools/assemble-thesis.py'),
                    '--body', str(output), '--output', str(args.output)], check=True)
    subprocess.run([sys.executable, str(root / 'tools/render-thesis.py'),
                    '--docx', str(args.output)], check=True)
else:
    print(output)
