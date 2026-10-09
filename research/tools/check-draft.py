#!/usr/bin/env python3
"""Validate current manuscript delivery, citations, Word structure and local provenance.

Run after export. --require-local-sources also checks retained source hashes;
those originals are intentionally absent from the public GitHub checkout.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--require-local-sources', action='store_true')
args = parser.parse_args()
errors = []

def require(condition, message):
    if not condition:
        errors.append(message)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def rows(path):
    with path.open(encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))

def citation_ids(text):
    result = set()
    for bracket in re.findall(r'\[([^\]\n]+)\]', text):
        for part in bracket.split(';'):
            match = re.match(r'\s*(\d+)(?:[–-](\d+))?', part)
            if match:
                a = int(match[1])
                b = int(match[2]) if match[2] else a
                result.update(range(a, b + 1))
    return result

def near(a, b, tolerance=3):
    return a is not None and abs(a - b) <= tolerance

md_path = ROOT / 'drafts/working-draft.md'
docx_path = ROOT / 'drafts/working-draft.docx'
text = md_path.read_text()
body, marker, bibliography = text.partition('# Список использованных источников')
require(bool(marker), 'Missing bibliography heading in Markdown')
references = [int(x) for x in re.findall(r'(?m)^(\d+)\. ', bibliography)]
require(references == list(range(1, len(references) + 1)), 'Reference numbering is not contiguous')
require(citation_ids(body) == set(references), 'Orphan or missing manuscript references')
sources = rows(ROOT / 'sources.csv')
evidence = rows(ROOT / 'evidence.csv')
mapping = rows(ROOT / 'literature/citation-map.csv')
source_ids = {s['source_id'] for s in sources}
require(len(source_ids) == len(sources), 'Duplicate source IDs')
require(len({r['claim_id'] for r in evidence}) == len(evidence), 'Duplicate claim IDs')
require({int(r['reference_number']) for r in mapping} == set(references), 'Citation map differs from bibliography')
require(len(mapping) == len(references), 'Duplicate citation-map numbers')
for row in mapping:
    require(row['source_id'] in source_ids, 'Unknown source in citation map: ' + row['source_id'])
for row in evidence:
    for sid in re.findall(r'S\d{3}', row['source_id']):
        require(sid in source_ids, f"Unknown source in {row['claim_id']}: {sid}")

manifest = json.loads((ROOT / 'literature/file-manifest.json').read_text())
require(source_ids <= {r.get('source_id') for r in manifest}, 'Source missing from provenance manifest')
local_missing = []
hashes_checked = 0
for row in manifest:
    path = ROOT / row['path']
    if not path.exists():
        local_missing.append(row['path'])
        continue
    hashes_checked += 1
    require(digest(path) == row['sha256'], 'Source hash changed: ' + row['path'])
if args.require_local_sources:
    require(not local_missing, 'Required retained sources missing: ' + ', '.join(local_missing))

preflights = list((ROOT / 'literature/pdf-preflight').glob('*.json'))
require(len(preflights) == 11, 'Expected eleven structural PDF preflights')
for path in preflights:
    record = json.loads(path.read_text())
    require(record['verdict'] == 'PASS', 'PDF preflight is not PASS: ' + path.name)
    pdf = ROOT.parent / record['file']
    if pdf.exists():
        require(digest(pdf) == record['sha256'], 'Stale PDF preflight: ' + path.name)

# Recheck the new course-backed evidence against the specified local pages.
course_sources = {s['source_id']: s for s in sources if s['source_type'] == 'user-provided course material'}
excerpt_checks = 0
for row in evidence:
    if row['source_id'] not in course_sources or not row['source_excerpt']:
        continue
    sid = row['source_id']
    source = course_sources[sid]
    section_num = re.search(r'section-(\d+)\.pdf', source['local_file'])[1]
    extracted = ROOT / f'literature/course/extracted/section-{section_num}.json'
    if not extracted.exists():
        continue
    pages = json.loads(extracted.read_text())['pages']
    match = re.search(r'с\.\s*(\d+)(?:–(\d+))?', row['page_or_locator'])
    require(bool(match), 'Missing page locator: ' + row['claim_id'])
    if not match:
        continue
    first, last = int(match[1]), int(match[2] or match[1])
    require(1 <= first <= last <= len(pages), 'Invalid pages: ' + row['claim_id'])
    selected = ' '.join(pages[first - 1:last])
    norm = lambda s: re.sub(r'\s+', ' ', s).casefold().strip()
    require(norm(row['source_excerpt']) in norm(selected), 'Excerpt absent from selected pages: ' + row['claim_id'])
    excerpt_checks += 1

with zipfile.ZipFile(docx_path) as archive:
    require(archive.testzip() is None, 'Corrupt DOCX container')
doc = Document(docx_path)
word_text = '\n'.join(''.join(p.itertext(qn('w:t'))) for p in doc._element.iter(qn('w:p')))
require(citation_ids(word_text) == citation_ids(body), 'Word and Markdown citations differ')
expected_headings = re.findall(r'(?m)^## (.+)$', body)
actual_headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
require(actual_headings == expected_headings, 'Word and Markdown subsection headings differ')
for heading in actual_headings:
    require(not heading.endswith('.'), 'Terminal dot in heading: ' + heading)
chapter_headings = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
require(chapter_headings[0].paragraph_format.page_break_before is False, 'Initial chapter starts with forced break')
for paragraph in chapter_headings:
    require(paragraph.text == paragraph.text.upper(), 'Chapter heading is not uppercase')
for ref in references:
    require(any(p.text.startswith(f'{ref} ') for p in doc.paragraphs), f'Word reference {ref} missing or has a terminal dot')

for candidate in [docx_path, ROOT / 'plan.docx', ROOT / 'drafts/chapter-1.docx']:
    with zipfile.ZipFile(candidate) as archive:
        require(archive.testzip() is None, 'Corrupt Word file: ' + candidate.name)
    output = Document(candidate)
    for section in output.sections:
        checks = [
            (section.page_width, Cm(21), 'A4 width'),
            (section.page_height, Cm(29.7), 'A4 height'),
            (section.left_margin, Cm(3), 'left margin'),
            (section.right_margin, Cm(1.5), 'right margin'),
            (section.top_margin, Cm(2), 'top margin'),
            (section.bottom_margin, Cm(2), 'bottom margin'),
        ]
        for actual, expected, label in checks:
            require(near(actual, expected, 1000), candidate.name + ': ' + label)
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        style = output.styles[name]
        require(style.font.name == 'Times New Roman', candidate.name + ': body font')
        require(near(style.font.size, Pt(14)), candidate.name + ': body font size')
        require(style.paragraph_format.line_spacing == 1.5, candidate.name + ': body line spacing')
        require(near(style.paragraph_format.first_line_indent, Cm(1.25), 1000), candidate.name + ': paragraph indent')
    require(any('PAGE' in element.get(qn('w:instr'), '') for section in output.sections
                for element in section.footer._element.iter(qn('w:fldSimple'))), candidate.name + ': missing PAGE field')
    for table in output.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    require(paragraph.paragraph_format.line_spacing == 1, candidate.name + ': table line spacing')
                    for run in paragraph.runs:
                        require(run.font.name == 'Times New Roman' and near(run.font.size, Pt(12)), candidate.name + ': table font')

for table in doc.tables:
    require(bool(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')), 'Missing repeated table header')
    require(abs(sum(c.width for c in table.columns) - Cm(16.5)) < 2000, 'Table exceeds body width')

result = {
    'status': 'failed' if errors else 'passed', 'body_words': len(body.split()),
    'subsections': [s.split()[0] for s in expected_headings],
    'references': len(references), 'sources': len(sources),
    'evidence_entries': len(evidence), 'tables': len(doc.tables),
    'pdf_preflights': len(preflights), 'local_source_hashes_checked': hashes_checked,
    'course_excerpts_checked': excerpt_checks, 'local_sources_unavailable': local_missing,
    'artifact_hashes': {str(p.relative_to(ROOT)): digest(p) for p in [md_path, docx_path, ROOT / 'plan.md', ROOT / 'plan.docx']},
    'checked': 'Word container and basic structural formatting; current subsection and citation correspondence; source identities, local source hashes and new course excerpt anchors when available',
    'not_checked': ['Visual pagination and line count after tables', 'Complete final thesis and department-specific acceptance', 'Unwritten comparative results', 'Publication authorship and date of course PDFs', 'Primary Russian rules, ship certificates and independent trial protocols'],
    'errors': errors,
}
dest = ROOT / 'build/working-draft-checks.json'
dest.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(result, ensure_ascii=False, indent=2))
raise SystemExit(1 if errors else 0)
